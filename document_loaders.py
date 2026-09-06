"""
Loaders for non-PDF document formats (docx, xlsx, pptx, txt, md, csv, tsv,
html, rtf), plus the dispatcher that routes every supported file - PDF or
not - to the right loader.

Every loader returns the same shape the PDF pipeline uses in
heading_detection.py:

    {page_number_0_indexed: [(section, text), ...]}

so ingestion.py can treat docx/xlsx/txt files identically to PDFs
downstream (chunking, metadata stamping, filters).

Two structural differences from PDFs are handled here:

- Pagination: these formats have no page breaks, so text is cut into
  pseudo-pages of ~PSEUDO_PAGE_CHARS characters (config.py). A page filter
  like "page 3" still works - it just means the third ~3000-character
  segment, not a physical sheet.
- Section labels: each format contributes whatever structural signals it
  actually has (docx heading styles, xlsx sheet names, pptx slide numbers,
  markdown '#' headings); everything else is "General Section", matching
  the PDF pipeline's graceful-degradation convention.
"""
import os
import re

from config import PSEUDO_PAGE_CHARS, SUPPORTED_DOC_EXTENSIONS

# =====================================================================
# SHARED HELPERS
# =====================================================================
def _paginate(section, chunks, pages, page):
    """Append `chunks` (an iterable of text pieces) to `pages` under
    `section`, cutting into pseudo-pages of PSEUDO_PAGE_CHARS chars.
    Returns the next free page number."""
    buf, chars = [], 0
    for chunk in chunks:
        if not chunk:
            continue
        buf.append(chunk)
        chars += len(chunk)
        if chars >= PSEUDO_PAGE_CHARS:
            pages.setdefault(page, []).append((section, "\n".join(buf)))
            page += 1
            buf, chars = [], 0
    if buf:
        pages.setdefault(page, []).append((section, "\n".join(buf)))
        page += 1
    return page


def _read_text(path):
    """Read a file as text with encoding sniffing (UTF-8 BOM, UTF-8,
    UTF-16 only when a BOM is present, then cp1252 as a last resort).
    UTF-16 is BOM-gated deliberately: a plain cp1252/latin-1 file of even
    byte length would otherwise "successfully" decode as UTF-16 into
    mojibake, silently corrupting the indexed text."""
    with open(path, "rb") as f:
        raw = f.read()
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return raw.decode("utf-16")
        except UnicodeDecodeError:
            pass
    for enc in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return raw.decode("utf-8", errors="replace")


# =====================================================================
# DOCX
# =====================================================================
_HEADING_STYLE_RE = re.compile(r"^(heading|title|subtitle)", re.IGNORECASE)


def load_docx(path):
    """docx paragraphs grouped under their current heading style. Tables
    are appended as 'cell | cell' rows (python-docx doesn't interleave
    them with paragraphs, so they trail the paragraph stream)."""
    import docx

    document = docx.Document(path)

    pages = {}
    page = 0
    section = "General Section"
    buf, chars = [], 0

    def flush():
        nonlocal buf, chars, page
        if buf:
            pages.setdefault(page, []).append((section, "\n".join(buf)))
            page += 1
            buf, chars = [], 0

    for para in document.paragraphs:
        style = (para.style.name or "") if para.style else ""
        text = para.text.strip()
        if not text:
            continue
        if _HEADING_STYLE_RE.match(style):
            flush()
            section = text
            buf.append(text)  # the heading line itself stays in the content
            chars = len(text)
            continue
        buf.append(text)
        chars += len(text)
        if chars >= PSEUDO_PAGE_CHARS:
            flush()
    flush()

    rows = []
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                rows.append(" | ".join(cells))
    if rows:
        page = _paginate(section, rows, pages, page)
    return pages


# =====================================================================
# XLSX
# =====================================================================
def load_xlsx(path):
    """One (sheet, rows) stream per worksheet; rows become 'cell | cell'
    lines, the sheet title becomes the section label."""
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        pages = {}
        page = 0
        for ws in workbook.worksheets:
            section = f"Sheet: {ws.title}"
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    rows.append(" | ".join(cells))
            page = _paginate(section, rows, pages, page)
        return pages
    finally:
        workbook.close()


# =====================================================================
# PPTX
# =====================================================================
def load_pptx(path):
    """One page per slide; the section label is the slide number."""
    from pptx import Presentation

    prs = Presentation(path)
    pages = {}
    for i, slide in enumerate(prs.slides):
        texts = [
            shape.text.strip()
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and (shape.text or "").strip()
        ]
        pages[i] = [(f"Slide {i + 1}", "\n".join(texts))]
    return pages


# =====================================================================
# PLAIN-TEXT FAMILY (txt / md / csv / tsv / html / htm / rtf)
# =====================================================================
_TAG_RE = re.compile(r"<[^>]+>")
_RTF_CONTROL_RE = re.compile(r"\\[a-zA-Z]+-?\d* ?")

_MD_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")


def load_text_document(path):
    """Markdown gets '#' headings as section labels; every other plain
    text format (txt, csv, tsv, html, rtf) is one 'General Section'
    stream. HTML tags and RTF control words are stripped best-effort."""
    text = _read_text(path)
    ext = os.path.splitext(path)[1].lower()

    if ext in (".html", ".htm"):
        text = _TAG_RE.sub(" ", text)
    elif ext == ".rtf":
        text = _RTF_CONTROL_RE.sub(" ", text)

    pages = {}
    if ext == ".md":
        page = 0
        section = "General Section"
        buf, chars = [], 0
        for line in text.split("\n"):
            m = _MD_HEADING_RE.match(line.strip())
            if m:
                if buf:
                    pages.setdefault(page, []).append((section, "\n".join(buf)))
                    page += 1
                    buf, chars = [], 0
                section = m.group(2).strip()
                buf.append(line.strip())
                chars = len(line)
                continue
            if line.strip():
                buf.append(line)
                chars += len(line)
                if chars >= PSEUDO_PAGE_CHARS:
                    pages.setdefault(page, []).append((section, "\n".join(buf)))
                    page += 1
                    buf, chars = [], 0
        if buf:
            pages.setdefault(page, []).append((section, "\n".join(buf)))
        return pages

    # Everything else: one stream, cut into pseudo-pages (preferring to
    # break at newlines near the boundary where possible).
    pages = {}
    page = 0
    pos = 0
    length = len(text)
    while pos < length:
        end = min(pos + PSEUDO_PAGE_CHARS, length)
        if end < length:
            nl = text.rfind("\n", pos, end)
            if nl > pos + PSEUDO_PAGE_CHARS // 2:
                end = nl
        piece = text[pos:end].strip()
        if piece:
            pages[page] = [("General Section", piece)]
        page += 1
        pos = end
    return pages


# =====================================================================
# DISPATCHER
# =====================================================================
def supported_extension(path):
    return os.path.splitext(path)[1].lower() in SUPPORTED_DOC_EXTENSIONS


def load_document_pages(path):
    """Route one file to its loader. Returns {page: [(section, text)]} -
    the same shape as heading_detection.build_page_texts_with_sections."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        from heading_detection import build_page_texts_with_sections
        return build_page_texts_with_sections(path)
    if ext == ".docx":
        return load_docx(path)
    if ext == ".xlsx":
        return load_xlsx(path)
    if ext == ".pptx":
        return load_pptx(path)
    if ext in (".txt", ".md", ".csv", ".tsv", ".html", ".htm", ".rtf"):
        return load_text_document(path)
    raise ValueError(f"Unsupported file type: {ext}")
